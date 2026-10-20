"""
Monthly SZoV reports uploaded as they come from the call-centre system (docs/SZOV_REPORTS.md).

Team reports (``report_YYYY-MM*.xlsx``) come for the line (sheets «Отработанные часы», «Звонки»,
«Эффективность») and for chats («Отработанные часы», «Чаты», «Среднее время ответа»). Both may carry
«Тренинги», «Тех. сбои» and «Офлайн активность» (hours that count towards the norm), «Штрафы»
(lateness) and the operator's «Норма часов (ч)» for the month. «др.» marks a day the operator spent
in another team. The QA report (``monthly_report_dates_YYYY-MM.xlsx``) gives reviewers' scores per
day and each operator's average for the month.

Operators are matched by full name. What the reports say is kept by operator, day and kind
(``ReportDayFact``), so a Monday–Sunday week that crosses two months is built from both monthly
reports, and a month exported before its end waits for the rest. Every reported day goes to
analytics; a week goes to the contest once each of its days is reported for the operator. Weeks are
never closed here, so no coins are awarded.
"""

from __future__ import annotations

import io
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, timedelta
from zipfile import BadZipFile, ZipFile

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import managed_operators_filter
from app.core.errors import DomainError
from app.models.contest import (
    ContestWeek,
    MetricDefinition,
    OperatorDayMetric,
    OperatorWeekMetric,
    ReportDayFact,
    ReportMonth,
)
from app.models.enums import Channel, WeekStatus
from app.models.user import User
from app.services import weekly as weekly_service
from app.services.economy import local_today
from app.services.months import MONTHS

MAX_FILES = 12
MAX_FILE_BYTES = 5 * 1024 * 1024
HOURS, CALLS, EFFECTIVE = "Отработанные часы", "Звонки", "Эффективность"
CHATS, RESPONSE = "Чаты", "Среднее время ответа"
#: Hours that count towards the norm besides the hours on the line or in chats.
EXTRA = ("Тренинги", "Тех. сбои", "Офлайн активность")
FINES = "Штрафы"
#: The operator was in another team that day: no data and no norm here.
AWAY = "др."
WEEK_HOURS, DAY_HOURS = 40, 8
LINE, CHAT = str(Channel.LINE), str(Channel.CHAT)
#: The source of the day and week values written from the reports.
SOURCE = "report"
#: Per channel: hours against the norm, hours above it, and the reviewers' score.
CODES = {
    LINE: {"norm": "hours_norm", "overtime": "overtime", "quality": "quality"},
    CHAT: {"norm": "chat_hours_norm", "overtime": "chat_overtime", "quality": "chat_quality"},
}
_DAY = re.compile(r"^(\d{2})\.(\d{2})$")
_MONTH = re.compile(r"(20\d{2})-(0[1-9]|1[0-2])")
#: Kinds of facts: a team's report tells where the operator was that day, the reviewers' — scores.
TEAMS = (LINE, CHAT)
QUALITY = "quality"
#: Someone with hours in the last week of a report carries the month's norm to the month's end.
ACTIVE_DAYS = 7
#: A month that is over is reported to its end when its last hours are this close to it.
END_DAYS = 3
#: What a report row of the day gives besides the hours, by the fact's key.
DAY_KINDS = ("calls", "effective", "chats", "response")


def normalize_name(name: object) -> str:
    return re.sub(r"\s+", " ", str(name or "")).strip().lower().replace("ё", "е")


def _number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return float(value)
    text = re.sub(r"[\s  ]", "", str(value)).replace(",", ".")
    try:
        return float(text) if text else None
    except ValueError:
        return None


def _away(value: object) -> bool:
    return isinstance(value, str) and value.strip().lower() == AWAY


def _day_columns(header: tuple, year: int, month: int) -> dict[int, date]:
    columns = {}
    for index, title in enumerate(header):
        match = _DAY.match(str(title or "").strip())
        if match and int(match.group(2)) == month:
            try:
                columns[index] = date(year, month, int(match.group(1)))
            except ValueError:
                continue
    return columns


def _columns(header: tuple) -> dict[str, int]:
    return {str(title).strip().casefold(): i for i, title in enumerate(header) if title is not None}


def _cell(row: tuple, column: int | None) -> object:
    return row[column] if column is not None and column < len(row) else None


def month_days(year: int, month: int) -> list[date]:
    days = [date(year, month, 1) + timedelta(days=i) for i in range(31)]
    return [day for day in days if day.month == month]


def touching_weeks(year: int, month: int) -> list[list[date]]:
    """Monday–Sunday weeks with at least one day in the month, the edge ones included."""
    days = month_days(year, month)
    monday = days[0] - timedelta(days=days[0].weekday())
    weeks = []
    while monday <= days[-1]:
        weeks.append([monday + timedelta(days=i) for i in range(7)])
        monday += timedelta(days=7)
    return weeks


@dataclass
class Team:
    """One operator in the reports of one channel (the line or chats)."""

    #: days[kind][day]: «hours» worked (the report counts trainings, tech failures and offline
    #: activity in them), «extra» (those three), «calls», «effective», «chats», «response»
    #: (seconds), «fines» (tenge).
    days: dict = field(default_factory=lambda: defaultdict(dict))
    #: Days in this team: «Отработанные часы» is not «др.».
    present: set = field(default_factory=set)
    #: Days marked «др.» in «Отработанные часы».
    away: set = field(default_factory=set)
    rate: float | None = None
    #: The month's «Норма часов (ч)», summed over the channel's reports; None without the column.
    norm: float | None = None
    fines: bool = False
    #: From the fines summary: lateness count and the count of other reasons.
    late: int | None = None
    other_fines: int = 0

    def hours(self, day: date) -> float:
        return self.days["hours"].get(day, 0.0)

    def is_away(self, day: date) -> bool:
        """In another team that day: «др.» in each of the channel's reports with the operator."""
        return day in self.away and day not in self.present

    def span(self, days: list[date], last: date) -> set[date]:
        """Days that carry the norm: from the first day with hours to the last one, without «др.».
        Zeros before someone started or after they left carry none; someone with hours in the
        last week of the report carries it to the month's end — the days not reported yet
        included, so a month exported before its end shares the norm as the whole month will."""
        worked = sorted(day for day, hours in self.days["hours"].items() if hours > 0)
        if not worked:
            return set()
        stop = days[-1] if worked[-1] >= last - timedelta(days=ACTIVE_DAYS) else worked[-1]
        return {day for day in days if worked[0] <= day <= stop and not self.is_away(day)}

    def norm_shares(self, days: list[date], last: date) -> dict[date, float]:
        """The report's norm shared evenly by the calendar days of the span; without the norm
        column (older reports) — 40 h × rate a week."""
        span = self.span(days, last)
        if not span:
            return {}
        share = self.norm / len(span) if self.norm else WEEK_HOURS * (self.rate or 1) / 7
        return dict.fromkeys(span, share)

    def lateness(self, day: date) -> int:
        return 1 if self.days["fines"].get(day, 0) > 0 else 0

    @property
    def lateness_known(self) -> bool:
        """Each day with a fine is a lateness only when the month has no other reasons."""
        return self.fines and self.late is not None and not (self.late and self.other_fines)

    def fact(self, day: date, norm: float) -> dict:
        """What the report says about the operator's day, as kept in ``ReportDayFact``."""
        if self.is_away(day):
            return {"away": True}
        data = {"hours": self.hours(day), "norm": norm, "rate": self.rate or 1}
        if extra := self.days["extra"].get(day):
            data["extra"] = extra
        for kind in DAY_KINDS:
            if day in self.days[kind]:
                data[kind] = self.days[kind][day]
        if fine := self.days["fines"].get(day):
            data["fine"] = fine
        if self.lateness_known:
            data["late"] = self.lateness(day)
        return data


@dataclass
class ParsedReports:
    year: int
    month: int
    files: list[dict] = field(default_factory=list)
    #: teams[(channel, name)] — what the line or chat reports say about the operator.
    teams: dict = field(default_factory=dict)
    #: quality[name][day] = reviewers' scores.
    quality: dict = field(default_factory=lambda: defaultdict(lambda: defaultdict(list)))
    #: The QA report's average for the month and how many scores it has (dated or not).
    quality_month: dict[str, tuple[float, int]] = field(default_factory=dict)
    #: Full names as written in the reports, for people missing on the site.
    labels: dict[str, str] = field(default_factory=dict)

    @property
    def names(self) -> set[str]:
        return {name for _, name in self.teams} | set(self.quality) | set(self.quality_month)

    def team(self, channel: str, name: str) -> Team:
        return self.teams.setdefault((channel, name), Team())


def _open(filename: str, content: bytes):
    if len(content) > MAX_FILE_BYTES:
        raise DomainError(f"{filename}: файл больше 5 МБ", code="file_too_large")
    if not filename.lower().endswith(".xlsx"):
        raise DomainError(f"{filename}: нужен файл XLSX", code="unsupported_file")
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                raise DomainError(f"{filename}: файл слишком большой", code="file_too_large")
        from openpyxl import load_workbook

        return load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, KeyError, ValueError, OSError) as exc:
        raise DomainError(f"{filename}: не удалось прочитать XLSX", code="invalid_file") from exc


def _person(row: tuple, parsed: ParsedReports) -> str:
    """The operator of a row; «Итого» and empty rows are nobody."""
    name = normalize_name(row[0] if row else None)
    if not name or name == "итого":
        return ""
    parsed.labels.setdefault(name, re.sub(r"\s+", " ", str(row[0])).strip())
    return name


def _read_days(book, title: str, parsed: ParsedReports, channel: str, kind: str) -> None:
    """A sheet of one value per operator per day; hours add up, an average is kept as it is."""
    if title not in book.sheetnames:
        return
    rows = book[title].iter_rows(values_only=True)
    days = _day_columns(next(rows, ()), parsed.year, parsed.month)
    seen = set()
    for row in rows:
        name = _person(row, parsed)
        if not name or name in seen:
            continue
        seen.add(name)
        values = parsed.team(channel, name).days[kind]
        for column, day in days.items():
            value = _number(_cell(row, column))
            if value is not None:
                values[day] = value if kind == "response" else values.get(day, 0.0) + value


def _read_hours(book, parsed: ParsedReports, channel: str) -> set[str]:
    rows = book[HOURS].iter_rows(values_only=True)
    header = next(rows, ())
    days, columns = _day_columns(header, parsed.year, parsed.month), _columns(header)
    rate = columns.get("ставка")
    norm = next((i for title, i in columns.items() if title.startswith("норма часов")), None)
    people = set()
    for row in rows:
        name = _person(row, parsed)
        if not name or name in people:
            continue
        people.add(name)
        team = parsed.team(channel, name)
        if value := _number(_cell(row, rate)):
            team.rate = value
        if norm is not None and (value := _number(_cell(row, norm))) is not None:
            team.norm = (team.norm or 0.0) + value
        for column, day in days.items():
            cell = _cell(row, column)
            if _away(cell):
                team.away.add(day)
                continue
            team.present.add(day)
            if (value := _number(cell)) is not None:
                team.days["hours"][day] = team.days["hours"].get(day, 0.0) + value
    return people


def _read_fines(book, parsed: ParsedReports, channel: str) -> None:
    """«Штрафы»: tenge per day, then the summary by reason (lateness = minutes × 50)."""
    if FINES not in book.sheetnames:
        return
    rows = list(book[FINES].iter_rows(values_only=True))
    if not rows:
        return
    days = _day_columns(rows[0], parsed.year, parsed.month)
    end = len(rows)
    for index, row in enumerate(rows[1:], start=1):
        first = str(_cell(row, 0) or "").strip()
        if first.casefold() == "итого" or first.startswith("Легенда"):
            end = index
            break
        name = _person(row, parsed)
        if not name:
            continue
        team = parsed.team(channel, name)
        team.fines = True
        for column, day in days.items():
            value = _number(_cell(row, column))
            if value:
                team.days["fines"][day] = team.days["fines"].get(day, 0.0) + value
    header = next(
        (i for i in range(end, len(rows)) if str(_cell(rows[i], 0) or "").strip() == "ФИО"),
        None,
    )
    if header is None:
        return
    columns = _columns(rows[header])
    late = columns.get("опоздание, кол-во")
    others = [
        i for title, i in columns.items()
        if title.endswith(", кол-во") and not title.startswith("опоздание")
    ]
    for row in rows[header + 1:]:
        name = _person(row, parsed)
        if not name:
            if str(_cell(row, 0) or "").strip():
                break
            continue
        team = parsed.team(channel, name)
        team.late = int(_number(_cell(row, late)) or 0) if late is not None else None
        team.other_fines = int(sum(_number(_cell(row, i)) or 0 for i in others))


def _read_team(book, parsed: ParsedReports, channel: str) -> int:
    people = _read_hours(book, parsed, channel)
    sheets = (
        {CALLS: "calls", EFFECTIVE: "effective"} if channel == LINE
        else {CHATS: "chats", RESPONSE: "response"}
    )
    for title, kind in sheets.items():
        _read_days(book, title, parsed, channel, kind)
    for title in EXTRA:
        _read_days(book, title, parsed, channel, "extra")
    _read_fines(book, parsed, channel)
    return len(people)


def _read_quality(book, parsed: ParsedReports) -> int:
    people = set()
    for sheet in book.worksheets:
        section, days, columns = "", {}, {}
        for row in sheet.iter_rows(values_only=True):
            first = str((row[0] if row else None) or "").strip()
            if first.startswith("Оценки"):
                section = first
                continue
            if first == "ФИО":
                days, columns = _day_columns(row, parsed.year, parsed.month), _columns(row)
                continue
            if not first or not section.startswith("Оценки по датам"):
                continue
            name = normalize_name(first)
            for column, day in days.items():
                for part in str(_cell(row, column) or "").split(","):
                    value = _number(part)
                    if value is not None and part.strip():
                        parsed.quality[name][day].append(value)
                        people.add(name)
            average = _number(_cell(row, columns.get("средний балл")))
            count = _number(_cell(row, columns.get("кол-во")))
            if average is not None and count:
                parsed.quality_month[name] = (average, int(count))
                people.add(name)
            if name in people:
                parsed.labels.setdefault(name, first)
    return len(people)


def detect_month(filenames: list[str], month: str | None) -> tuple[int, int]:
    if month:
        match = _MONTH.fullmatch(month.strip())
        if not match:
            raise DomainError("Месяц укажите в виде ГГГГ-ММ", code="invalid_month")
        return int(match.group(1)), int(match.group(2))
    found = {m.groups() for name in filenames if (m := _MONTH.search(name))}
    if len(found) == 1:
        year, number = found.pop()
        return int(year), int(number)
    if len(found) > 1:
        raise DomainError("Файлы относятся к разным месяцам", code="mixed_months")
    raise DomainError(
        "Не удалось определить месяц по названию файлов — выберите его вручную",
        code="month_required",
    )


def parse(files: list[tuple[str, bytes]], month: str | None = None) -> ParsedReports:
    if not files:
        raise DomainError("Выберите файлы отчётов", code="no_files")
    if len(files) > MAX_FILES:
        raise DomainError(f"Можно загрузить до {MAX_FILES} файлов", code="too_many_files")
    year, number = detect_month([name for name, _ in files], month)
    parsed = ParsedReports(year, number)
    for filename, content in files:
        book = _open(filename, content)
        try:
            sheets = set(book.sheetnames)
            if HOURS in sheets and {CALLS, EFFECTIVE} <= sheets:
                kind, people = LINE, _read_team(book, parsed, LINE)
            elif HOURS in sheets and (CHATS in sheets or RESPONSE in sheets):
                kind, people = CHAT, _read_team(book, parsed, CHAT)
            else:
                people = _read_quality(book, parsed)
                kind = "quality" if people else "unknown"
        finally:
            book.close()
        parsed.files.append({"filename": filename, "kind": kind, "people": people})
    if not any(item["kind"] != "unknown" for item in parsed.files):
        raise DomainError(
            "Не найдено ни отчёта команды (линия или чаты), ни отчёта проверяющих",
            code="unknown_reports",
        )
    return parsed


def report_until(parsed: ParsedReports, today: date) -> dict[str, date]:
    """The last day each channel's team reports cover. A month that is over is reported to its
    end (its last days without hours, at most three, are days off); the current month — to the
    last day with hours, but not today: its shifts are still going on."""
    days = month_days(parsed.year, parsed.month)
    until = {}
    for channel in TEAMS:
        worked = [
            day
            for (team_channel, _), team in parsed.teams.items()
            if team_channel == channel
            for day, hours in team.days["hours"].items()
            if hours > 0
        ]
        if not worked:
            continue
        last = max(worked)
        if days[-1] < today:
            last = days[-1] if (days[-1] - last).days <= END_DAYS else last
        else:
            last = min(last, today - timedelta(days=1))
        if last >= days[0]:
            until[channel] = last
    return until


def month_facts(
    parsed: ParsedReports, ids: dict[str, int], until: dict[str, date]
) -> tuple[dict, set]:
    """What the uploaded reports say by (operator, day, kind), and which (operator, kind) pairs
    they speak for: a team report — up to its last day, the reviewers' — for the whole month."""
    days = month_days(parsed.year, parsed.month)
    facts, kinds = {}, set()
    for (channel, name), team in parsed.teams.items():
        user_id, last = ids.get(name), until.get(channel)
        if user_id is None or last is None:
            continue
        kinds.add((user_id, channel))
        shares = team.norm_shares(days, last)
        for day in days:
            if day <= last:
                facts[(user_id, day, channel)] = team.fact(day, shares.get(day, 0.0))
    for name in set(parsed.quality) | set(parsed.quality_month):
        user_id = ids.get(name)
        if user_id is None:
            continue
        kinds.add((user_id, QUALITY))
        for day, scores in parsed.quality.get(name, {}).items():
            if scores:
                facts[(user_id, day, QUALITY)] = {"scores": list(scores)}
    return facts, kinds


def merge_facts(stored: dict, new: dict, kinds: set, days: list[date]) -> dict:
    """The facts after the upload. A team report replaces what was kept about its days, except
    the days it marks «др.»: those keep what another team's report said. The reviewers' report
    replaces the operator's scores for the month."""
    merged = dict(stored)
    for user_id, kind in kinds:
        for day in days:
            key = (user_id, day, kind)
            fresh = new.get(key)
            if kind == QUALITY:
                if fresh is None:
                    merged.pop(key, None)
                else:
                    merged[key] = fresh
            elif fresh is not None:
                old = stored.get(key)
                if not (fresh.get("away") and old and not old.get("away")):
                    merged[key] = fresh
    return merged


def _team_values(channel: str, team: dict[date, dict], *, week: bool) -> dict[str, float]:
    values: dict[str, float] = {}
    codes = CODES[channel]
    facts = list(team.values())
    hours = sum(fact["hours"] for fact in facts)
    if hours > 0:
        norm = (
            sum(fact.get("norm", 0.0) for fact in facts) if week
            else DAY_HOURS * (facts[0].get("rate") or 1)
        )
        if norm > 0:
            values[codes["norm"]] = round(hours / norm * 100, 1)
            if week:
                values[codes["overtime"]] = round(max(0.0, hours - norm), 2)
    # Calls, efficiency and chats are over «База часов»: the hours on the line or in chats,
    # without the trainings, tech failures and offline activity.
    base = sum(max(0.0, fact["hours"] - fact.get("extra", 0.0)) for fact in facts)
    if channel == LINE:
        if base > 0:
            effective = sum(fact.get("effective", 0.0) for fact in facts)
            values["efficiency"] = round(min(100.0, effective / base * 100), 1)
            values["calls_per_hour"] = round(sum(f.get("calls", 0.0) for f in facts) / base, 2)
    else:
        # As the report: the plain average of the days' averages.
        times = [fact["response"] for fact in facts if "response" in fact]
        if times:
            values["chat_response_time"] = round(sum(times) / len(times), 1)
        if base > 0 and any("chats" in fact for fact in facts):
            values["chats_per_hour"] = round(sum(f.get("chats", 0.0) for f in facts) / base, 2)
    return values


def _lateness(teams: dict[str, dict[date, dict]], *, week: bool) -> int | None:
    """A day with a fine is a lateness when the month's fines tell it (no other reasons). Counted
    in the teams the operator worked in; unknown in any of them — unknown."""
    counted = [
        team for team in teams.values()
        if (
            sum(fact["hours"] for fact in team.values()) > 0 if week
            else any(fact["hours"] > 0 or fact.get("fine") for fact in team.values())
        )
    ]
    if not counted or any("late" not in fact for team in counted for fact in team.values()):
        return None
    return sum(fact["late"] for team in counted for fact in team.values())


def values_of(
    facts: dict[str, dict[date, dict]],
    days: list[date],
    channel: str,
    *,
    week: bool,
    month_quality: list | tuple | None = None,
) -> dict[str, float]:
    """One operator's values for a week or a day from what the reports say about its days."""
    teams = {
        team: {
            day: fact for day in days
            if (fact := facts.get(team, {}).get(day)) and not fact.get("away")
        }
        for team in TEAMS
    }
    values: dict[str, float] = {}
    for team, rows in teams.items():
        if rows:
            values.update(_team_values(team, rows, week=week))

    def channel_of(day: date) -> str:
        """Where the operator worked that day; the site's channel when the reports don't tell."""
        inside = [team for team in TEAMS if day in teams[team]]
        return inside[0] if len(inside) == 1 else channel

    scores = facts.get(QUALITY, {})
    for team in TEAMS:
        marks = [
            score for day in days if channel_of(day) == team
            for score in scores.get(day, {}).get("scores", ())
        ]
        code = CODES[team]["quality"]
        if marks:
            values[code] = round(sum(marks) / len(marks), 1)
        elif week and month_quality and sum(f["hours"] for f in teams[team].values()) > 0:
            # The QA report has scores without a date: a week without dated ones gets the
            # month's average, as long as the operator worked that week.
            values[code] = round(month_quality[0], 1)
    late = _lateness(teams, week=week)
    if late is not None:
        values["lateness"] = late
    return values


@dataclass
class Coverage:
    """Which days are reported for an operator: a team report says something about the day, or
    the month's report of the operator's channel covers it and the operator is not in it."""

    #: (year, month) → {channel: the last reported day}.
    until: dict
    #: (user_id, (year, month)) — the operator is in a team report of that month.
    members: set

    def covers(self, user_id: int, channel: str, facts: dict, day: date) -> bool:
        if any(day in facts.get(team, {}) for team in TEAMS):
            return True
        month = (day.year, day.month)
        if (user_id, month) in self.members:
            return False
        last = self.until.get(month, {}).get(channel)
        return last is not None and day <= last


def _waiting(gaps: dict[tuple[int, int], set[str]], until: dict) -> str:
    """Why a week waits: gaps[(year, month)] — the channels of the operators missing those days."""
    parts = []
    for (year, month), channels in sorted(gaps.items()):
        reported = until.get((year, month), {})
        lacking = channels - set(reported)
        name = MONTHS[month - 1]
        if not reported or lacking == {LINE, CHAT}:
            parts.append(f"отчёта за {name}")
        elif lacking:
            parts.append(f"отчёта {'линии' if LINE in lacking else 'чатов'} за {name}")
        else:
            last = min(reported[channel] for channel in channels)
            parts.append(f"выгрузки после {last:%d.%m}")
    return "ждёт " + " и ".join(dict.fromkeys(parts))


def _month_quality(averages: dict, user_id: int, days: list[date]) -> list | None:
    """The reviewers' month average for a week: of the month of its Thursday, else the other."""
    for day in (days[3], days[0], days[-1]):
        if value := averages.get((day.year, day.month), {}).get(user_id):
            return value
    return None


async def _operators(session: AsyncSession, actor: User) -> dict[str, User]:
    scope = await managed_operators_filter(session, actor)
    users = await session.scalars(select(User).where(scope, User.is_active.is_(True)))
    return {normalize_name(user.full_name): user for user in users}


async def _stored_facts(session: AsyncSession, users: set[int], first: date, last: date) -> dict:
    if not users:
        return {}
    rows = await session.scalars(
        select(ReportDayFact).where(
            ReportDayFact.user_id.in_(users), ReportDayFact.day.between(first, last)
        )
    )
    return {(row.user_id, row.day, row.kind): row.data for row in rows}


async def _members(session: AsyncSession, users: set[int], months: set[tuple[int, int]]) -> set:
    """Who is in a team report of the given months."""
    members = set()
    for year, month in months:
        days = month_days(year, month)
        found = await session.scalars(
            select(ReportDayFact.user_id).distinct().where(
                ReportDayFact.user_id.in_(users),
                ReportDayFact.day.between(days[0], days[-1]),
                ReportDayFact.kind.in_(TEAMS),
            )
        )
        members |= {(user_id, (year, month)) for user_id in found}
    return members


def _notes(
    parsed: ParsedReports, found: list[str], operators: dict[str, User], by_user: dict,
    coverage: Coverage,
) -> list[dict]:
    """What the person uploading should know: data the reports could not place."""
    month = (parsed.year, parsed.month)
    groups: dict[str, list[str]] = defaultdict(list)
    for name in found:
        user = operators[name]
        label, channel = user.full_name, str(user.channel or LINE)
        teams = {team for team in TEAMS if (team, name) in parsed.teams}
        if teams and channel not in teams:
            groups["channel"].append(label)
        quality = parsed.quality_month.get(name)
        dated = sum(len(scores) for scores in parsed.quality.get(name, {}).values())
        if (user.id, month) not in coverage.members and (quality or dated):
            reported = channel in coverage.until.get(month, {})
            groups["not_in_team" if reported else "no_team"].append(label)
        elif quality and quality[1] > dated:
            groups["undated"].append(label)
        for team in (parsed.teams[(c, name)] for c in teams):
            if team.fines and team.late and team.other_fines:
                groups["lateness_unknown"].append(label)
            elif team.lateness_known and team.late and team.late > sum(
                map(team.lateness, team.present)
            ):
                groups["lateness_merged"].append(label)
    titles = {
        "channel": "В отчётах другое направление, чем на сайте: баллы считаются по направлению "
        "из карточки сотрудника",
        "undated": "Есть оценки без даты: в недели без оценок по датам записан средний балл за "
        "месяц",
        "no_team": "Отчёта команды за месяц ещё нет: оценки сохранены, недели запишутся, когда "
        "придёт отчёт команды",
        "not_in_team": "Нет в отчётах команд за месяц: в недели записано только качество",
        "lateness_unknown": "В штрафах есть и другие причины: опоздания по дням не определить, "
        "они не записаны",
        "lateness_merged": "Несколько опозданий в один день: засчитано по одному за день",
    }
    return [
        {"code": code, "title": titles[code], "names": sorted(set(groups[code]))}
        for code in titles
        if groups.get(code)
    ]


async def build_plan(session: AsyncSession, actor: User, parsed: ParsedReports) -> dict:
    """What the upload changes: the kept facts, the days of the month and the weeks touching it
    that are reported for their operators; the other weeks wait (``pending``)."""
    operators = await _operators(session, actor)
    found = sorted(name for name in parsed.names if name in operators)
    ids = {name: operators[name].id for name in found}
    channels = {operators[name].id: str(operators[name].channel or LINE) for name in found}
    month = (parsed.year, parsed.month)
    days, weeks = month_days(*month), touching_weeks(*month)
    until = report_until(parsed, local_today())
    stored = await _stored_facts(session, set(ids.values()), weeks[0][0], weeks[-1][-1])
    new, kinds = month_facts(parsed, ids, until)
    facts = merge_facts(stored, new, kinds, days)
    by_user = {user_id: defaultdict(dict) for user_id in ids.values()}
    for (user_id, day, kind), data in facts.items():
        by_user[user_id][kind][day] = data

    neighbours = {(day.year, day.month) for day in (weeks[0][0], weeks[-1][-1])} - {month}
    reports = {
        (item.year, item.month): item
        for item in await session.scalars(select(ReportMonth))
        if (item.year, item.month) in neighbours | {month}
    }
    reported = {
        key: {team: last for team in TEAMS if (last := getattr(item, f"{team}_until"))}
        for key, item in reports.items()
    }
    for team, last in until.items():
        current = reported.setdefault(month, {}).get(team)
        reported[month][team] = max(last, current) if current else last
    coverage = Coverage(reported, await _members(session, set(ids.values()), neighbours))
    for user_id, user_facts in by_user.items():
        team_days = (day for team in TEAMS for day in user_facts.get(team, {}))
        if any((day.year, day.month) == month for day in team_days):
            coverage.members.add((user_id, month))

    averages = {
        key: {int(user_id): value for user_id, value in (item.quality or {}).items()}
        for key, item in reports.items()
    }
    quality = {
        ids[name]: parsed.quality_month.get(name)
        for name in found
        if name in parsed.quality or name in parsed.quality_month
    }
    for user_id, value in quality.items():
        if value:
            averages.setdefault(month, {})[user_id] = list(value)
        else:
            averages.get(month, {}).pop(user_id, None)

    daily = {}
    for user_id, user_facts in by_user.items():
        for day in days:
            for code, value in values_of(user_facts, [day], channels[user_id], week=False).items():
                daily[(user_id, day, code)] = value
    planned, pending = [], []
    for week in weeks:
        rows, gaps, waiting = {}, defaultdict(set), 0
        for user_id, user_facts in by_user.items():
            channel = channels[user_id]
            missing = [
                day for day in week if not coverage.covers(user_id, channel, user_facts, day)
            ]
            if not missing:
                rows[user_id] = values_of(
                    user_facts, week, channel, week=True,
                    month_quality=_month_quality(averages, user_id, week),
                )
            elif any(day in user_facts.get(team, {}) for team in TEAMS for day in week):
                waiting += 1
                for day in missing:
                    gaps[(day.year, day.month)].add(channel)
        if any(rows.values()):
            planned.append({"days": week, "values": rows})
        if waiting:
            pending.append({"days": week, "operators": waiting, "reason": _waiting(gaps, reported)})

    unmatched = sorted(
        parsed.labels.get(name, name)
        for name in parsed.names
        if name not in operators and AWAY not in name
    )
    changed = {
        key: facts.get(key)
        for key in set(facts) | set(stored)
        if facts.get(key) != stored.get(key)
    }
    return {
        "found": found,
        "unmatched": unmatched,
        "without_data": sorted(set(operators) - set(found)),
        "names": {name: operators[name].full_name for name in found},
        "month": month,
        "users": sorted(ids.values()),
        "until": until,
        "facts": changed,
        "quality": quality,
        "weeks": planned,
        "pending": pending,
        "daily": daily,
        "notes": _notes(parsed, found, operators, by_user, coverage),
    }


async def ensure_metrics(session: AsyncSession, parsed: ParsedReports | None = None) -> None:
    """Hours are judged against the report's norm, so the catalogue needs «Выполнение нормы
    часов»; a chat report brings the chats' metrics. Missing ones are created as the site's
    defaults (``app/db/init_db.py``)."""
    from app.db.init_db import CHAT_METRICS, CHAT_QUALITY, DEFAULT_METRICS

    wanted = [item for item in DEFAULT_METRICS if item["code"] == "hours_norm"]
    if parsed is not None and any(item["kind"] == CHAT for item in parsed.files):
        wanted += [CHAT_QUALITY, *CHAT_METRICS]
    existing = {
        metric.code: metric
        for metric in await session.scalars(
            select(MetricDefinition).where(
                MetricDefinition.code.in_([item["code"] for item in wanted])
            )
        )
    }
    for item in wanted:
        metric = existing.get(item["code"])
        if metric is None:
            session.add(MetricDefinition(**{"is_active": True, **item}))
        elif item["code"] == "hours_norm" and not metric.is_active:
            metric.is_active = True
    await session.flush()


async def _save_facts(session: AsyncSession, plan: dict) -> None:
    changes = plan["facts"]
    if not changes:
        return
    days = [day for _, day, _ in changes]
    rows = {
        (row.user_id, row.day, row.kind): row
        for row in await session.scalars(
            select(ReportDayFact).where(
                ReportDayFact.user_id.in_({user_id for user_id, _, _ in changes}),
                ReportDayFact.day.between(min(days), max(days)),
            )
        )
    }
    for (user_id, day, kind), data in changes.items():
        row = rows.get((user_id, day, kind))
        if data is None:
            if row is not None:
                await session.delete(row)
        elif row is None:
            session.add(ReportDayFact(user_id=user_id, day=day, kind=kind, data=data))
        else:
            row.data = data
    await session.flush()


async def _save_month(session: AsyncSession, plan: dict) -> None:
    if not plan["until"] and not plan["quality"]:
        return
    year, month = plan["month"]
    item = await session.scalar(
        select(ReportMonth).where(ReportMonth.year == year, ReportMonth.month == month)
    )
    if item is None:
        item = ReportMonth(year=year, month=month, quality={})
        session.add(item)
    for team, last in plan["until"].items():
        current = getattr(item, f"{team}_until")
        if current is None or last > current:
            setattr(item, f"{team}_until", last)
    if plan["quality"]:
        quality = dict(item.quality or {})
        for user_id, value in plan["quality"].items():
            if value:
                quality[str(user_id)] = list(value)
            else:
                quality.pop(str(user_id), None)
        item.quality = quality
    await session.flush()


async def save(session: AsyncSession, actor: User, plan: dict, *, recalculate: bool) -> dict:
    await _save_facts(session, plan)
    await _save_month(session, plan)
    daily, days = plan["daily"], month_days(*plan["month"])
    existing = (
        {
            (row.user_id, row.day, row.metric_code): row
            for row in await session.scalars(
                select(OperatorDayMetric).where(
                    OperatorDayMetric.user_id.in_(plan["users"]),
                    OperatorDayMetric.day.between(days[0], days[-1]),
                )
            )
        }
        if plan["users"]
        else {}
    )
    for key, row in existing.items():
        # A day the reports no longer give a value for loses the one they gave before.
        if key not in daily and row.source == SOURCE:
            await session.delete(row)
    for (user_id, day, code), value in daily.items():
        row = existing.get((user_id, day, code))
        if row is None:
            session.add(
                OperatorDayMetric(
                    user_id=user_id, day=day, metric_code=code, value=value, source=SOURCE
                )
            )
        else:
            row.value, row.source = value, SOURCE

    weeks = []
    for item in plan["weeks"]:
        week: ContestWeek = await weekly_service.get_or_create_week(session, item["days"][0])
        if week.status == WeekStatus.CLOSED:
            weeks.append({"label": week.label, "status": "closed", "operators": 0})
            continue
        current = {
            (row.user_id, row.metric_code): row
            for row in await session.scalars(
                select(OperatorWeekMetric).where(OperatorWeekMetric.week_id == week.id)
            )
        }
        for (user_id, code), row in current.items():
            # The operator's week is built anew: what the reports no longer give goes.
            values = item["values"].get(user_id)
            if values is not None and row.source == SOURCE and code not in values:
                await session.delete(row)
        for user_id, values in item["values"].items():
            for code, value in values.items():
                row = current.get((user_id, code))
                if row is None:
                    session.add(
                        OperatorWeekMetric(
                            week_id=week.id,
                            user_id=user_id,
                            metric_code=code,
                            value=value,
                            source=SOURCE,
                        )
                    )
                else:
                    row.value, row.source = value, SOURCE
        await session.flush()
        await weekly_service.invalidate_calculation(session, week)
        operators = sum(1 for values in item["values"].values() if values)
        if recalculate and operators:
            await weekly_service.calculate_week(session, week)
        weeks.append({"label": week.label, "status": str(week.status), "operators": operators})
    return {"days": len(daily), "weeks": weeks}
