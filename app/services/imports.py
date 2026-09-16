"""Bounded CSV/XLSX parsing and shared validation; preview never writes data."""

from __future__ import annotations

import csv
import io
import math
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import managed_operators_filter
from app.core.errors import DomainError, PermissionDeniedError
from app.models.contest import ContestWeek, MetricDefinition
from app.models.enums import Role
from app.models.user import User
from app.schemas.admin import MetricValueIn
from app.schemas.imports import ImportIssue, ImportPreviewOut

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000
MAX_COLUMNS = 128
MAX_VALUES = 100000

_ALIASES = {
    "operator_id": "user_id",
    "id": "user_id",
    "логин": "login",
    "operator_login": "login",
    "показатель": "metric_code",
    "значение": "value",
    "дата": "date",
}
_IDENTITY_COLUMNS = {"user_id", "login", "date"}


async def validate_values(session: AsyncSession, actor: User, values: list[MetricValueIn]) -> None:
    """Validate the apply request again; callers cannot bypass file preview."""
    users = {
        user.id: user
        for user in await session.scalars(
            select(User).where(User.id.in_({value.user_id for value in values}))
        )
    }
    visible_ids = set(
        await session.scalars(select(User.id).where(await managed_operators_filter(session, actor)))
    )
    codes = set(
        await session.scalars(
            select(MetricDefinition.code).where(MetricDefinition.is_active.is_(True))
        )
    )
    seen: set[tuple[int, str]] = set()
    for value in values:
        user = users.get(value.user_id)
        if user is None or user.role != Role.OPERATOR or not user.is_active:
            raise DomainError(
                "Показатели можно загрузить только для действующего оператора", code="unknown_users"
            )
        if user.id not in visible_ids:
            raise PermissionDeniedError("Оператор не входит в вашу зону ответственности")
        if value.metric_code not in codes:
            raise DomainError(
                f"Неизвестный или отключённый показатель: {value.metric_code}",
                code="unknown_metric",
            )
        if not math.isfinite(value.value):
            raise DomainError("Значение должно быть конечным числом", code="invalid_value")
        key = (value.user_id, value.metric_code)
        if key in seen:
            raise DomainError("Повторяется показатель одного оператора", code="duplicate_metric")
        seen.add(key)


def _read_rows(filename: str, content: bytes) -> list[list[object]]:
    if len(content) > MAX_FILE_BYTES:
        raise DomainError("Файл превышает 5 МБ", code="file_too_large")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        try:
            body = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                body = content.decode("cp1251")
            except UnicodeDecodeError as exc:
                raise DomainError(
                    "Не удалось прочитать кодировку CSV", code="invalid_file"
                ) from exc
        try:
            dialect = csv.Sniffer().sniff(body[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        iterator = csv.reader(io.StringIO(body, newline=""), dialect, strict=True)
        workbook = None
    elif suffix == ".xlsx":
        try:
            # Limit decompressed input as well as the uploaded ZIP container.
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                    raise DomainError("Распакованный XLSX слишком большой", code="file_too_large")
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            if not workbook.worksheets:
                raise DomainError("В книге нет листов", code="invalid_file")
            sheet = workbook.worksheets[0]
            if (sheet.max_column or 0) > MAX_COLUMNS or (sheet.max_row or 0) > MAX_ROWS + 1:
                workbook.close()
                raise DomainError(
                    "Допускается до 10 000 строк и 128 столбцов", code="file_too_large"
                )
            iterator = sheet.iter_rows(values_only=True)
        except ImportError as exc:
            raise DomainError(
                "Поддержка XLSX недоступна. Используйте CSV", code="xlsx_unavailable"
            ) from exc
        except (BadZipFile, KeyError, ValueError, OSError) as exc:
            raise DomainError("Не удалось прочитать XLSX", code="invalid_file") from exc
    else:
        raise DomainError("Поддерживаются файлы CSV и XLSX", code="unsupported_file")

    rows: list[list[object]] = []
    try:
        for row in iterator:
            if len(rows) > MAX_ROWS or len(row) > MAX_COLUMNS:
                raise DomainError(
                    "Допускается до 10 000 строк и 128 столбцов", code="file_too_large"
                )
            rows.append(list(row))
    except (csv.Error, ValueError, TypeError) as exc:
        raise DomainError("Повреждённая таблица", code="invalid_file") from exc
    finally:
        if workbook is not None:
            workbook.close()
    return rows


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _number(value: object) -> float:
    if isinstance(value, bool):
        raise ValueError("Boolean is not a metric")
    result = float(_text(value).replace("\u00a0", "").replace(" ", "").replace(",", "."))
    if not math.isfinite(result):
        raise ValueError("Non-finite metric")
    return result


async def preview_file(
    session: AsyncSession,
    actor: User,
    week: ContestWeek,
    filename: str,
    content: bytes,
) -> ImportPreviewOut:
    result = ImportPreviewOut(filename=filename)
    try:
        rows = _read_rows(filename, content)
    except DomainError as exc:
        result.errors.append(ImportIssue(code=exc.code, message=exc.message))
        return result

    if not rows:
        result.errors.append(ImportIssue(code="empty_file", message="Файл пуст"))
        return result
    headers = [_ALIASES.get(_text(value).lower(), _text(value).lower()) for value in rows[0]]
    if not any(column in headers for column in ("user_id", "login")):
        result.errors.append(
            ImportIssue(row=1, code="missing_identity", message="Нужен столбец user_id или login")
        )
    if any(not column for column in headers) or len(headers) != len(set(headers)):
        result.errors.append(
            ImportIssue(
                row=1,
                code="invalid_headers",
                message="Заголовки должны быть заполнены и не повторяться",
            )
        )
    long_format = "metric_code" in headers or "value" in headers
    result.format = "long" if long_format else "wide"
    if long_format and not {"metric_code", "value"}.issubset(headers):
        result.errors.append(
            ImportIssue(
                row=1, code="missing_columns", message="Нужны оба столбца metric_code и value"
            )
        )
    codes = set(
        await session.scalars(
            select(MetricDefinition.code).where(MetricDefinition.is_active.is_(True))
        )
    )
    data_columns = [column for column in headers if column not in _IDENTITY_COLUMNS]
    unknown = set(data_columns) - ({"metric_code", "value"} if long_format else codes)
    for column in sorted(unknown):
        result.errors.append(
            ImportIssue(
                row=1,
                field=column,
                code="unknown_metric",
                message=f"Неизвестный столбец или отключённый показатель: {column}",
            )
        )
    if not data_columns:
        result.errors.append(
            ImportIssue(row=1, code="missing_columns", message="Нет столбцов показателей")
        )
    if result.errors:
        return result

    users = list(
        await session.scalars(
            select(User).where(
                await managed_operators_filter(session, actor),
                User.role == Role.OPERATOR,
                User.is_active.is_(True),
            )
        )
    )
    by_id = {user.id: user for user in users}
    by_login = {user.login: user for user in users}
    seen: set[tuple[int, str]] = set()
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(_text(cell) for cell in row):
            continue
        result.total_rows += 1
        if len(row) > len(headers):
            result.errors.append(
                ImportIssue(
                    row=row_number,
                    code="invalid_row",
                    message="В строке больше ячеек, чем заголовков",
                )
            )
            continue
        record = dict(zip(headers, row + [None] * (len(headers) - len(row)), strict=True))
        user = None
        identity = _text(record.get("user_id"))
        if identity:
            try:
                number = _number(record["user_id"])
                if not number.is_integer() or number <= 0:
                    raise ValueError("Invalid id")
                user = by_id.get(int(number))
            except (ValueError, OverflowError):
                pass
        login = _text(record.get("login"))
        if login:
            login_user = by_login.get(login)
            if identity and (user is None or login_user is None or user.id != login_user.id):
                result.errors.append(
                    ImportIssue(
                        row=row_number,
                        field="login",
                        code="identity_mismatch",
                        message="user_id и login не соответствуют одному доступному оператору",
                    )
                )
                continue
            user = login_user
        if user is None:
            result.errors.append(
                ImportIssue(
                    row=row_number,
                    field="user_id" if identity else "login",
                    code="unknown_operator",
                    message="Действующий оператор не найден в вашей зоне ответственности",
                )
            )
            continue
        if _text(record.get("date")):
            try:
                raw_date = record["date"]
                if isinstance(raw_date, datetime):
                    parsed_date = raw_date.date()
                elif isinstance(raw_date, date):
                    parsed_date = raw_date
                else:
                    value = _text(raw_date)
                    parsed_date = (
                        datetime.strptime(value, "%d.%m.%Y").date()
                        if "." in value
                        else date.fromisoformat(value)
                    )
                if not week.starts_on <= parsed_date <= week.ends_on:
                    raise ValueError("Outside selected period")
            except (ValueError, TypeError):
                result.errors.append(
                    ImportIssue(
                        row=row_number,
                        field="date",
                        code="invalid_date",
                        message="Дата должна входить в выбранную неделю",
                    )
                )
                continue
        cells = (
            [(_text(record["metric_code"]), record["value"])]
            if long_format
            else [(column, record[column]) for column in data_columns]
        )
        for code, raw_value in cells:
            if code not in codes:
                result.errors.append(
                    ImportIssue(
                        row=row_number,
                        field="metric_code",
                        code="unknown_metric",
                        message=f"Неизвестный или отключённый показатель: {code}",
                    )
                )
                continue
            if not _text(raw_value):
                result.warnings.append(
                    ImportIssue(
                        row=row_number,
                        field=code,
                        code="missing_value",
                        message="Нет данных: значение пропущено, не заменено нулём",
                    )
                )
                continue
            try:
                value = _number(raw_value)
            except (ValueError, TypeError, OverflowError):
                result.errors.append(
                    ImportIssue(
                        row=row_number,
                        field=code,
                        code="invalid_value",
                        message="Ожидается конечное число без формулы",
                    )
                )
                continue
            key = (user.id, code)
            if key in seen:
                result.errors.append(
                    ImportIssue(
                        row=row_number,
                        field=code,
                        code="duplicate",
                        message="Показатель этого оператора уже есть в файле",
                    )
                )
                continue
            seen.add(key)
            result.values.append(MetricValueIn(user_id=user.id, metric_code=code, value=value))
            if len(result.values) >= MAX_VALUES:
                result.errors.append(
                    ImportIssue(
                        code="too_many_values",
                        message="Файл содержит слишком много значений; разделите его на части",
                    )
                )
                break
        if len(result.values) >= MAX_VALUES:
            break
    result.valid_values = len(result.values)
    result.operator_count = len({value.user_id for value in result.values})
    if not result.values and not result.errors:
        result.errors.append(
            ImportIssue(code="empty_data", message="В файле нет заполненных показателей")
        )
    result.can_apply = bool(result.values) and not result.errors
    return result
