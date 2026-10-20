import io
from datetime import date, timedelta

from openpyxl import Workbook
from sqlalchemy import select

from app.models.contest import (
    ContestWeek,
    MetricDefinition,
    OperatorDayMetric,
    OperatorWeekMetric,
    ReportMonth,
)
from app.models.user import User
from app.services import szov_reports
from tests.conftest import auth, login

AUGUST = [f"{day:02d}.08" for day in range(1, 32)]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def team_report(name: str, rate: float = 1) -> bytes:
    """The older team report: no norm column, so a week is judged against 40 h × rate."""
    book = Workbook()
    book.remove(book.active)
    for title, value in (("Отработанные часы", 8), ("Звонки", 40), ("Эффективность", 6)):
        sheet = book.create_sheet(title)
        sheet.append(["Оператор", "Ставка", *AUGUST])
        # Weekends are empty: 5 working days × 8 h = 40 h a week.
        days = [value if (day + 4) % 7 < 5 else None for day in range(1, 32)]
        sheet.append([name, rate, *days])
        sheet.append(["Другой Человек др.", 1, *days])
        sheet.append(["Итого", None, *days])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def quality_report(name: str) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Проверяющий"
    sheet.append(["Оценки по датам · СЗоВ"])
    sheet.append(["ФИО", *AUGUST])
    sheet.append([name, None, None, "90, 100", "80"])
    sheet.append(["Неизвестный Оператор", None, None, "70"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def upload(files: dict[str, bytes]):
    return [("files", (name, content, XLSX)) for name, content in files.items()]


async def test_reports_preview_then_save_days_and_weeks(client, session, operator, head):
    headers = auth(await login(client, head.login))
    files = {
        "report_2026-08_3.xlsx": team_report(operator.full_name, rate=0.5),
        "monthly_report_dates_2026-08.xlsx": quality_report(operator.full_name),
    }
    preview = await client.post(
        "/api/v1/admin/day-metrics/reports/preview", files=upload(files), headers=headers
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["month"] == "2026-08"
    assert [f["kind"] for f in body["files"]] == ["line", "quality"]
    assert body["matched_names"] == [operator.full_name]
    assert body["unmatched"] == ["Неизвестный Оператор"]
    assert [w["starts_on"] for w in body["weeks"]] == [
        "2026-08-03",
        "2026-08-10",
        "2026-08-17",
        "2026-08-24",
    ]
    assert (await session.scalar(select(OperatorDayMetric.id))) is None

    saved = await client.post(
        "/api/v1/admin/day-metrics/reports", files=upload(files), headers=headers
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["detail"].startswith("Сохранено")
    days = {
        (row.day.isoformat(), row.metric_code): row.value
        for row in await session.scalars(
            select(OperatorDayMetric).where(OperatorDayMetric.user_id == operator.id)
        )
    }
    # Half rate: 8 h against a 4 h day norm.
    assert days[("2026-08-03", "hours_norm")] == 200
    assert days[("2026-08-03", "quality")] == 95
    assert days[("2026-08-03", "calls_per_hour")] == 5
    assert days[("2026-08-03", "efficiency")] == 75
    week = await session.scalar(select(ContestWeek).where(ContestWeek.iso_week == 32))
    weekly = {
        row.metric_code: row.value
        for row in await session.scalars(
            select(OperatorWeekMetric).where(OperatorWeekMetric.week_id == week.id)
        )
    }
    assert weekly["hours_norm"] == 200 and weekly["quality"] == 90
    await session.refresh(week)
    assert week.status == "calculated"

    again = await client.post(
        "/api/v1/admin/day-metrics/reports", files=upload(files), headers=headers
    )
    assert again.status_code == 200
    count = len(list(await session.scalars(select(OperatorDayMetric))))
    assert count == len(days)


async def test_reports_need_a_month_and_a_known_file(client, operator, supervisor):
    headers = auth(await login(client, supervisor.login))
    nameless = {"team.xlsx": team_report(operator.full_name)}
    missing = await client.post(
        "/api/v1/admin/day-metrics/reports/preview", files=upload(nameless), headers=headers
    )
    assert missing.status_code >= 400
    assert "месяц" in missing.json()["detail"]
    chosen = await client.post(
        "/api/v1/admin/day-metrics/reports/preview",
        files=upload(nameless),
        data={"month": "2026-08"},
        headers=headers,
    )
    assert chosen.status_code == 200 and chosen.json()["matched"] == 1

    blank = io.BytesIO()
    Workbook().save(blank)
    unknown = await client.post(
        "/api/v1/admin/day-metrics/reports/preview",
        files=upload({"other_2026-08.xlsx": blank.getvalue()}),
        headers=headers,
    )
    assert unknown.status_code >= 400

    # A supervisor saves values; the head recalculates the week later.
    saved = await client.post(
        "/api/v1/admin/day-metrics/reports",
        files=upload({"report_2026-08_1.xlsx": team_report(operator.full_name)}),
        headers=headers,
    )
    assert saved.status_code == 200, saved.text
    assert {w["status"] for w in saved.json()["weeks"]} == {"open"}


async def test_operator_cannot_upload_reports(client, operator):
    headers = auth(await login(client, operator.login))
    files = {"report_2026-08_3.xlsx": team_report(operator.full_name)}
    response = await client.post(
        "/api/v1/admin/day-metrics/reports/preview", files=upload(files), headers=headers
    )
    assert response.status_code == 403


SEPTEMBER = [f"{day:02d}.09" for day in range(1, 31)]


def weekday(day: int) -> bool:
    # 1 September 2026 is a Tuesday.
    return (day + 0) % 7 not in (5, 6)


def book_of(sheets: dict[str, list[list]]) -> bytes:
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def team_sheet(name: str, rate: float, norm: float, per_day, *, total: str = "Итого"):
    head = ["Оператор", "Ставка", "Норма часов (ч)", *SEPTEMBER, total]
    return [head, [name, rate, norm, *[per_day(day) for day in range(1, 31)]]]


def extra_sheet(name: str, per_day):
    return [["Оператор", *SEPTEMBER, "Всего (ч)"], [name, *[per_day(day) for day in range(1, 31)]]]


def fines_sheet(name: str, rate: float, fines: dict[int, str], late: int, no_show: int = 0):
    return [
        ["ФИО", "Ставка", *SEPTEMBER, "Итого"],
        [name, rate, *[fines.get(day) for day in range(1, 31)]],
        ["Итого"],
        [],
        ["Легенда причин"],
        ["Причина", "Расчет"],
        ["Опоздание", "минуты x 50"],
        [],
        ["Сводка по операторам"],
        ["ФИО", "Ставка", "Опоздание, кол-во", "Опоздание, минуты", "Опоздание, сумма",
         "Корп такси, кол-во", "Корп такси, сумма", "Не выход, кол-во", "Не выход, сумма",
         "Всего штрафов", "Итого"],
        [name, rate, late, late * 10, late * 500, 0, 0, no_show, no_show * 10000],
        ["Итого"],
    ]


async def test_line_chat_and_reviewers_reports_as_exported(client, session, operator, head):
    """September 2026 as the call-centre system exports it: a line team, a chat team with «др.»
    days, reviewers' scores with and without a date, trainings and fines."""
    chat = User(login="chat1", full_name="Чатов Оператор", role="operator", channel="chat",
                hashed_password="x")
    session.add(chat)
    await session.commit()
    line_name, chat_name = operator.full_name, chat.full_name

    def line_hours(day):  # 8 h every weekday, 2 of them a training on Monday 7 September
        return 8 if weekday(day) else 0

    line = book_of({
        "Отработанные часы": team_sheet(line_name, 1, 168, line_hours, total="Итого часов"),
        "Перерыв": team_sheet(line_name, 1, 168, lambda day: 0.5),
        "Звонки": team_sheet(line_name, 1, 168, lambda day: 10 if weekday(day) else 0),
        "Эффективность": team_sheet(line_name, 1, 168, lambda day: 6 if weekday(day) else 0),
        "Тренинги": extra_sheet(line_name, lambda day: 2 if day == 7 else None),
        "Тех. сбои": extra_sheet(line_name, lambda day: None),
        "Штрафы": fines_sheet(line_name, 1, {8: "500", 9: "250"}, late=2),
    })

    def chat_hours(day):  # in another team until 6 September, then 4 h every weekday
        return "др." if day < 7 else 4 if weekday(day) else 0

    chats = book_of({
        "Отработанные часы": team_sheet(chat_name, 0.5, 84, chat_hours, total="Итого часов"),
        "Чаты": team_sheet(chat_name, 0.5, 84, lambda day: "др." if day < 7 else 50),
        "Средняя оценка": team_sheet(chat_name, 0.5, 84, lambda day: 4.9),
        "Среднее время ответа": team_sheet(
            chat_name, 0.5, 84,
            lambda day: "др." if day < 7 else (100 if day % 7 in (0, 1, 2) else 130)
            if weekday(day) else None,
        ),
        "Офлайн активность": extra_sheet(chat_name, lambda day: "др." if day < 7 else None),
        "Штрафы": fines_sheet(chat_name, 0.5, {8: "10 000", 9: "650"}, late=1, no_show=1),
    })
    quality = book_of({
        "Проверяющий": [
            ["Оценки по датам · СЗоВ — Служба заботы о водителях"],
            ["ФИО", *SEPTEMBER, "Средний балл", "Кол-во", "План"],
            [line_name, *[{8: "90, 100", 15: "80"}.get(day) for day in range(1, 31)], 88, 4, 4],
            [chat_name, *[None] * 30, 92, 3, 3],
            ["Неизвестный Оператор", *[None] * 30, 70, 1, 1],
            [],
            ["Оценки проверяющего 169 · СЗоВ"],
            ["ФИО", *SEPTEMBER, "Средний балл", "Кол-во", "План"],
            [line_name, *[None] * 30, 0, 0, 4],
        ],
    })
    files = {
        "report_2026-09_1.xlsx": line,
        "report_2026-09.xlsx": chats,
        "monthly_report_dates_2026-09.xlsx": quality,
    }
    headers = auth(await login(client, head.login))
    preview = await client.post(
        "/api/v1/admin/day-metrics/reports/preview", files=upload(files), headers=headers
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert [f["kind"] for f in body["files"]] == ["line", "chat", "quality"]
    assert body["matched"] == 2 and body["unmatched"] == ["Неизвестный Оператор"]
    notes = {note["code"]: note["names"] for note in body["notes"]}
    assert notes["undated"] == sorted([line_name, chat_name])
    assert notes["lateness_unknown"] == [chat_name]

    saved = await client.post(
        "/api/v1/admin/day-metrics/reports", files=upload(files), headers=headers
    )
    assert saved.status_code == 200, saved.text
    weeks = {
        week.iso_week: week.id
        for week in await session.scalars(select(ContestWeek).where(ContestWeek.iso_year == 2026))
    }
    assert sorted(weeks) == [37, 38, 39]
    weekly = {
        (row.week_id, row.user_id, row.metric_code): row.value
        for row in await session.scalars(select(OperatorWeekMetric))
    }

    def week(number, user, code):
        return weekly.get((weeks[number], user.id, code))

    # The line: 40 h (2 of them a training) against 168 h × 7 / 30 days = 39.2 h.
    assert week(37, operator, "hours_norm") == 102.0 and week(37, operator, "overtime") == 0.8
    # Calls and efficiency are over the 38 h on the line, without the training.
    assert week(37, operator, "calls_per_hour") == 1.32
    assert week(37, operator, "efficiency") == 78.9
    assert week(37, operator, "quality") == 95 and week(38, operator, "quality") == 80
    # No dated score in the last week: the month's average, the operator worked.
    assert week(39, operator, "quality") == 88
    assert week(37, operator, "lateness") == 2 and week(38, operator, "lateness") == 0
    # Chats: «др.» until 6 September, so 84 h is shared by the 24 days from the 7th.
    assert week(37, chat, "chat_hours_norm") == 81.6 and week(37, chat, "chat_overtime") == 0
    assert week(37, chat, "chat_response_time") == 112
    # «Чаты» over the 20 h in chats: 50 a day all week.
    assert week(37, chat, "chats_per_hour") == 17.5
    assert week(37, chat, "chat_quality") == 92 and week(37, chat, "quality") is None
    # A no-show among the fines: lateness by day is unknown and not written.
    assert week(37, chat, "lateness") is None
    assert week(37, chat, "hours_norm") is None and week(37, chat, "calls_per_hour") is None

    days = {
        (row.user_id, row.day.isoformat(), row.metric_code): row.value
        for row in await session.scalars(select(OperatorDayMetric))
    }
    assert days[(operator.id, "2026-09-08", "quality")] == 95
    assert days[(operator.id, "2026-09-08", "lateness")] == 1
    assert days[(operator.id, "2026-09-10", "lateness")] == 0
    assert days[(operator.id, "2026-09-07", "hours_norm")] == 100
    assert days[(chat.id, "2026-09-07", "chat_hours_norm")] == 100
    assert (chat.id, "2026-09-03", "chat_hours_norm") not in days
    codes = set(
        await session.scalars(
            select(MetricDefinition.code).where(MetricDefinition.channel == "chat")
        )
    )
    assert {"chat_hours_norm", "chat_overtime", "chat_response_time", "chats_per_hour"} <= codes


# Reports of any month, kept by day: weeks across months and months exported before their end.


def month_of(year: int, month: int) -> list[date]:
    days = [date(year, month, 1) + timedelta(days=n) for n in range(31)]
    return [day for day in days if day.month == month]


def weekdays(day: date) -> bool:
    return day.weekday() < 5


def line_book(year: int, month: int, rows: list[tuple], *, late: dict | None = None) -> bytes:
    """A line team's report: rows of (name, norm, hours by day); 10 calls and 6 effective hours a
    worked day; «Штрафы» with the lateness by day when given."""
    days = month_of(year, month)
    head = ["Оператор", "Ставка", "Норма часов (ч)", *[d.strftime("%d.%m") for d in days]]

    def sheet(value):
        return [head] + [
            [name, 1, norm, *[value(hours(d)) for d in days]] for name, norm, hours in rows
        ]

    def worked(cell):
        return isinstance(cell, int | float) and cell > 0

    sheets = {
        "Отработанные часы": sheet(lambda cell: cell),
        "Звонки": sheet(lambda cell: 10 if worked(cell) else cell if cell == "др." else None),
        "Эффективность": sheet(lambda cell: 6 if worked(cell) else None),
    }
    if late is not None:
        fines = [["ФИО", "Ставка", *[d.strftime("%d.%m") for d in days], "Итого"]]
        fines += [[name, 1, *["500" if d in late.get(name, ()) else None for d in days]]
                  for name, _, _ in rows]
        fines += [["Итого"], [], ["Сводка по операторам"],
                  ["ФИО", "Ставка", "Опоздание, кол-во", "Не выход, кол-во"]]
        fines += [[name, 1, len(late.get(name, ())), 0] for name, _, _ in rows]
        sheets["Штрафы"] = fines
    return book_of(sheets)


async def _week_values(session, iso_week: int, user_id: int) -> dict[str, float]:
    week = await session.scalar(
        select(ContestWeek).where(ContestWeek.iso_year == 2026, ContestWeek.iso_week == iso_week)
    )
    if week is None:
        return {}
    return {
        row.metric_code: row.value
        for row in await session.scalars(
            select(OperatorWeekMetric).where(
                OperatorWeekMetric.week_id == week.id, OperatorWeekMetric.user_id == user_id
            )
        )
    }


def _post(client, headers, files, *, preview=False):
    path = "/api/v1/admin/day-metrics/reports" + ("/preview" if preview else "")
    return client.post(path, files=upload(files), headers=headers)


async def test_a_week_across_two_months_is_built_from_both_reports(
    client, session, operator, head, monkeypatch
):
    monkeypatch.setattr(szov_reports, "local_today", lambda: date(2026, 11, 5))
    headers = auth(await login(client, head.login))
    name = operator.full_name

    def hours(day):
        return 8 if weekdays(day) else 0

    september = {"report_2026-09_1.xlsx": line_book(2026, 9, [(name, 176, hours)])}
    preview = (await _post(client, headers, september, preview=True)).json()
    assert [w["label"] for w in preview["weeks"]] == ["2026-W37", "2026-W38", "2026-W39"]
    assert preview["reported_to"] == {"line": "2026-09-30"}
    # The edge weeks wait for the neighbouring months.
    assert [(w["label"], w["reason"]) for w in preview["pending"]] == [
        ("2026-W36", "ждёт отчёта за август"), ("2026-W40", "ждёт отчёта за октябрь"),
    ]
    assert (await _post(client, headers, september)).status_code == 200
    assert await _week_values(session, 40, operator.id) == {}

    october = {"report_2026-10_1.xlsx": line_book(2026, 10, [(name, 176, hours)])}
    saved = (await _post(client, headers, october)).json()
    assert "2026-W40" in [w["label"] for w in saved["weeks"]]
    assert [w["label"] for w in saved["pending"]] == ["2026-W44"]
    # 28.09–04.10: 40 h against 3 days of 176 h / 30 and 4 days of 176 h / 31 = 40.31 h.
    week = await _week_values(session, 40, operator.id)
    assert week["hours_norm"] == 99.2 and week["overtime"] == 0
    assert week["calls_per_hour"] == 1.25
    item = await session.scalar(select(ReportMonth).where(ReportMonth.month == 10))
    assert item.line_until == date(2026, 10, 31)


async def test_a_month_exported_before_its_end_waits_for_the_rest(
    client, session, operator, head, monkeypatch
):
    headers = auth(await login(client, head.login))
    name = operator.full_name
    late = {name: {date(2026, 10, 6)}}

    def hours_until(last):
        return lambda day: (8 if weekdays(day) else 0) if day.day <= last else None

    # Exported on Tuesday 20 October: the hours end on Monday the 19th.
    monkeypatch.setattr(szov_reports, "local_today", lambda: date(2026, 10, 20))
    partial = {"report_2026-10_1.xlsx": line_book(
        2026, 10, [(name, 176, hours_until(19))], late=late)}
    body = (await _post(client, headers, partial)).json()
    assert body["reported_to"] == {"line": "2026-10-19"}
    assert [w["label"] for w in body["weeks"]] == ["2026-W41", "2026-W42"]
    # The weeks the operator is in wait: for September and for the rest of the 19th's week.
    pending = {w["label"]: w["reason"] for w in body["pending"]}
    assert pending == {
        "2026-W40": "ждёт отчёта за сентябрь", "2026-W43": "ждёт выгрузки после 19.10",
    }
    # The month's norm is shared by all its 31 days, not by the 19 exported: 40 h of 39.74 h.
    week = await _week_values(session, 41, operator.id)
    assert week["hours_norm"] == 100.6 and week["lateness"] == 1
    # The week in progress gets nothing yet: no lateness 0, no hours.
    assert await _week_values(session, 43, operator.id) == {}
    days = {row.day for row in await session.scalars(
        select(OperatorDayMetric).where(OperatorDayMetric.user_id == operator.id))}
    assert max(days) == date(2026, 10, 19)

    # The whole month later: the rest of the weeks, the same norm for the ones written.
    monkeypatch.setattr(szov_reports, "local_today", lambda: date(2026, 11, 2))
    full = {"report_2026-10_1.xlsx": line_book(
        2026, 10, [(name, 176, hours_until(31))], late=late)}
    body = (await _post(client, headers, full)).json()
    assert body["reported_to"] == {"line": "2026-10-31"}
    assert [w["label"] for w in body["weeks"]] == ["2026-W41", "2026-W42", "2026-W43"]
    assert (await _week_values(session, 41, operator.id))["hours_norm"] == 100.6
    week = await _week_values(session, 43, operator.id)
    assert week["hours_norm"] == 100.6 and week["lateness"] == 0


async def test_a_move_between_teams_keeps_both_reports(
    client, session, operator, head, monkeypatch
):
    """Team reports uploaded one by one: the days one marks «др.» keep what the other said."""
    monkeypatch.setattr(szov_reports, "local_today", lambda: date(2026, 11, 5))
    headers = auth(await login(client, head.login))
    name = operator.full_name

    def first_team(day):
        return "др." if day.day > 15 else 8 if weekdays(day) else 0

    def second_team(day):
        return "др." if day.day <= 15 else 8 if weekdays(day) else 0

    for number, hours in ((1, first_team), (3, second_team)):
        files = {f"report_2026-09_{number}.xlsx": line_book(2026, 9, [(name, 88, hours)])}
        assert (await _post(client, headers, files)).status_code == 200
    # 14–20 September: Monday and Tuesday in the first team, the rest in the second.
    # Each team's 88 h is shared by its 15 days: 7 × 5.87 h = 41.07 h of norm for 40 h.
    week = await _week_values(session, 38, operator.id)
    assert week["hours_norm"] == 97.4 and week["calls_per_hour"] == 1.25
